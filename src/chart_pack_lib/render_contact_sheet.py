"""
Renders a static contact sheet from the ACTUAL saved .pptx (not a re-derivation from
deck_content.py) -- reads real shape geometry/fill/text straight out of the OOXML via
python-pptx, so this is a faithful proxy for each slide's final resting state.

No LibreOffice / PowerPoint is available in this environment to produce a literal
render, so this is the closest available substitute; PowerPoint-only effects (the
Morph/Wipe/Fade transitions and entrance-animation timing) are not depicted here --
only the shapes' final static positions, sizes, colors and text, which is what a
completed build looks like once every animation has played.
"""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "chart_pack" / "deck" / "QSR_LHI_Executive_Presentation.pptx"
OUT = ROOT / "chart_pack" / "deck" / "contact_sheet.png"

DPI = 110
FONT_DIR = Path(r"C:\Windows\Fonts")

_font_cache = {}


def get_font(name, size, bold=False, italic=False):
    key = (name, size, bold, italic)
    if key in _font_cache:
        return _font_cache[key]
    size_px = max(6, int(size * DPI / 72))
    candidates = []
    if name == "Cambria":
        candidates = ["cambriab.ttf"] if bold else ["cambria.ttc"]
    else:
        if bold and italic:
            candidates = ["calibriz.ttf"]
        elif bold:
            candidates = ["calibrib.ttf"]
        elif italic:
            candidates = ["calibrii.ttf"]
        else:
            candidates = ["calibri.ttf"]
    font = None
    for c in candidates:
        p = FONT_DIR / c
        if p.exists():
            try:
                font = ImageFont.truetype(str(p), size_px)
                break
            except Exception:
                pass
    if font is None:
        font = ImageFont.load_default()
    _font_cache[key] = font
    return font


def emu_to_px(v):
    return int(v / 914400 * DPI)


def draw_slide(slide, sw_px, sh_px):
    img = Image.new("RGB", (sw_px, sh_px), "white")
    draw = ImageDraw.Draw(img)
    for shape in slide.shapes:
        try:
            l, t, w, h = shape.left, shape.top, shape.width, shape.height
        except TypeError:
            continue
        if l is None:
            continue
        x0, y0 = emu_to_px(l), emu_to_px(t)
        x1, y1 = emu_to_px(l + (w or 0)), emu_to_px(t + (h or 0))

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                blob = shape.image.blob
                import io
                pic = Image.open(io.BytesIO(blob)).convert("RGB")
                pic = pic.resize((max(1, x1 - x0), max(1, y1 - y0)))
                img.paste(pic, (x0, y0))
            except Exception:
                pass
            continue

        is_connector = shape.shape_type is None and hasattr(shape, "line") and not shape.has_text_frame
        if is_connector or (hasattr(shape, "line") and not getattr(shape, "has_text_frame", False)):
            try:
                color = shape.line.color.rgb
                dash = None
                ln = shape.line._get_or_add_ln()
                pd = ln.find("{http://schemas.openxmlformats.org/drawingml/2006/main}prstDash")
                dash = pd is not None
                if x1 == x0:
                    line_pts = [(x0, y0), (x1, y1)]
                else:
                    line_pts = [(x0, y0), (x1, y1)]
                if dash:
                    _draw_dashed(draw, line_pts[0], line_pts[1], str(color), width=2)
                else:
                    draw.line(line_pts, fill=str(color), width=2)
            except Exception:
                pass
            continue

        try:
            fill_color = None
            if shape.fill.type is not None and shape.fill.type == 1:  # MSO_FILL.SOLID
                fill_color = "#" + str(shape.fill.fore_color.rgb)
            outline_color = None
            try:
                if shape.line.fill.type == 1:
                    outline_color = "#" + str(shape.line.color.rgb)
            except Exception:
                outline_color = None
            if fill_color or outline_color:
                draw.rectangle([x0, y0, max(x0, x1), max(y0, y1)], fill=fill_color,
                                outline=outline_color, width=3 if outline_color else 0)
        except Exception:
            pass

        if getattr(shape, "has_text_frame", False):
            tf = shape.text_frame
            box_w = max(1, x1 - x0)
            line_specs = []  # (text, font, color, align, size)
            for p in tf.paragraphs:
                for r in p.runs:
                    if not r.text:
                        continue
                    sz = r.font.size.pt if r.font.size else 12
                    fnt = get_font(r.font.name or "Calibri", sz, bool(r.font.bold), bool(r.font.italic))
                    col = "#" + str(r.font.color.rgb) if r.font.color and r.font.color.type is not None else "#141F28"
                    for wrapped in _wrap_text(draw, r.text, fnt, box_w):
                        line_specs.append((wrapped, fnt, col, p.alignment, sz))
            total_text_h = sum(int(sz * 1.22 * DPI / 72) for _, _, _, _, sz in line_specs)
            y_cursor = y0
            if tf.vertical_anchor is not None and str(tf.vertical_anchor).startswith("MIDDLE"):
                y_cursor = y0 + max(0, ((y1 - y0) - total_text_h) // 2)
            for text, fnt, col, align, sz in line_specs:
                bbox = draw.textbbox((0, 0), text, font=fnt)
                tw = bbox[2] - bbox[0]
                if align == PP_ALIGN.CENTER:
                    tx = x0 + max(0, ((x1 - x0) - tw) // 2)
                elif align == PP_ALIGN.RIGHT:
                    tx = x1 - tw
                else:
                    tx = x0
                draw.text((tx, y_cursor), text, font=fnt, fill=col)
                y_cursor += int(sz * 1.22 * DPI / 72)
    return img


def _wrap_text(draw, text, font, max_w_px):
    words = text.split(" ")
    lines, cur = [], ""
    for w in words:
        trial = (cur + " " + w).strip()
        tw = draw.textbbox((0, 0), trial, font=font)[2]
        if tw <= max_w_px or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines or [""]


def _draw_dashed(draw, p0, p1, color, width=2, dash=6, gap=5):
    import math
    x0, y0 = p0
    x1, y1 = p1
    dist = math.hypot(x1 - x0, y1 - y0)
    if dist == 0:
        return
    steps = int(dist // (dash + gap)) + 1
    for i in range(steps):
        s0 = i * (dash + gap)
        s1 = min(s0 + dash, dist)
        t0, t1 = s0 / dist, s1 / dist
        draw.line([(x0 + (x1 - x0) * t0, y0 + (y1 - y0) * t0),
                   (x0 + (x1 - x0) * t1, y0 + (y1 - y0) * t1)], fill=color, width=width)


def main():
    prs = Presentation(str(DECK))
    sw_px = emu_to_px(prs.slide_width)
    sh_px = emu_to_px(prs.slide_height)
    slides = list(prs.slides)
    # physical index (0-based) of each logical slide's FINAL resting state
    final_idx = [1, 3, 4, 6, 7, 8, 9, 10]
    labels = ["P1 Risk-band distribution", "P2 LEAKY vs SAFE", "P3 SHAP drivers",
              "P4 Aspect sentiment", "P5 High-risk share", "P6 LHI version agreement",
              "P7 Early warning", "Fig 11 Intervention priority"]

    panels = []
    for idx in final_idx:
        panels.append(draw_slide(slides[idx], sw_px, sh_px))

    scale = 0.42
    pw, ph = int(sw_px * scale), int(sh_px * scale)
    cols, rows = 2, 4
    pad = 22
    label_h = 28
    sheet_w = cols * pw + (cols + 1) * pad
    sheet_h = rows * (ph + label_h) + (rows + 1) * pad
    sheet = Image.new("RGB", (sheet_w, sheet_h), "#e7e4dd")
    d = ImageDraw.Draw(sheet)
    label_font = get_font("Calibri", 11, bold=True)
    for i, (panel, label) in enumerate(zip(panels, labels)):
        r, c = divmod(i, cols)
        x = pad + c * (pw + pad)
        y = pad + r * (ph + label_h + pad)
        thumb = panel.resize((pw, ph))
        d.rectangle([x - 2, y - 2, x + pw + 2, y + ph + 2], outline="#8b93a0", width=2)
        sheet.paste(thumb, (x, y))
        d.text((x, y + ph + 5), label, font=label_font, fill="#141F28")
    sheet.save(OUT)
    print("wrote", OUT, sheet.size)


if __name__ == "__main__":
    main()
